"""Manual / file-based import — the compliant acquisition path for portals that
cannot be automated (GeM bidplus and IREPS disallow automated access in
robots.txt; most portals put documents behind a CAPTCHA).

A business-development user drops a CSV/XLSX/YAML file (template in
``tests/fixtures/manual/manual_import_template.csv``) into the inbox, and
optionally downloaded tender documents into ``inbox/documents/<tender id>/``.

Also used by the ``fixture`` source type for offline demos and tests, where it
additionally accepts saved CPPP/GePNIC HTML pages.
"""

from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import yaml

from ..models import DocumentRef, ListingPage, OpportunityType, TenderRecord, TenderStatus, normalize_space, stable_hash
from ..timeutil import now_utc, parse_portal_datetime
from .base import DiscoveryContext, TenderSourceConnector, infer_opportunity_type, infer_state, parse_inr_amount

COLUMNS = [
    "platform", "tender_id", "ref_no", "title", "organisation", "department", "state", "opportunity_type",
    "category", "tender_category", "published_at", "closing_at", "opening_at", "estimated_value_inr", "emd_inr",
    "description", "official_url", "document_paths", "status", "evidence",
]


def _read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return [dict(r) for r in csv.DictReader(fh)]
    if suffix in (".yaml", ".yml"):
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        rows = data.get("records", data) if isinstance(data, dict) else data
        return [dict(r) for r in rows or []]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        rows = [{header[i]: v for i, v in enumerate(r) if i < len(header)} for r in it]
        wb.close()
        return [r for r in rows if any(v not in (None, "") for v in r.values())]
    return []


def row_to_record(row: dict[str, Any], source_id: str, default_platform: str, base_dir: Path, file_name: str) -> TenderRecord:
    def s(key: str) -> str | None:
        v = row.get(key)
        if v is None:
            return None
        v = normalize_space(str(v))
        return v or None

    title = s("title")
    if not title:
        raise ValueError("row has no title")
    tender_id = s("tender_id")
    org = s("organisation")
    published = _dt(row.get("published_at"))
    closing = _dt(row.get("closing_at"))
    docs: list[DocumentRef] = []
    for p in (s("document_paths") or "").split(";"):
        p = p.strip()
        if not p:
            continue
        if p.startswith(("http://", "https://")):
            docs.append(DocumentRef(url=p, filename=p.rsplit("/", 1)[-1], access="direct"))
        else:
            lp = (base_dir / p).resolve() if not Path(p).is_absolute() else Path(p)
            docs.append(DocumentRef(local_path=str(lp), filename=lp.name, access="manual"))
    otype_raw = s("opportunity_type")
    try:
        otype = OpportunityType(otype_raw) if otype_raw else infer_opportunity_type(title, s("ref_no"), tender_id)
    except ValueError:
        otype = infer_opportunity_type(otype_raw, title)
    status_raw = (s("status") or "active").lower()
    try:
        status = TenderStatus(status_raw)
    except ValueError:
        status = TenderStatus.UNKNOWN
    key = tender_id or f"h:{stable_hash(s('ref_no'), org, title, closing)}"
    return TenderRecord(
        source_id=source_id,
        platform=s("platform") or default_platform,
        source_record_key=key,
        source_tender_id=tender_id,
        ref_no=s("ref_no"),
        title=title,
        organisation=org,
        department=s("department"),
        state=s("state") or infer_state(org, title),
        opportunity_type=otype,
        category=s("category"),
        tender_category=s("tender_category"),
        published_at=published,
        closing_at=closing,
        opening_at=_dt(row.get("opening_at")),
        estimated_value_inr=_amount(row.get("estimated_value_inr")),
        emd_inr=_amount(row.get("emd_inr")),
        description=s("description"),
        official_url=s("official_url"),
        documents=docs,
        status=status,
        evidence=s("evidence") or f"Manual import file {file_name}",
        fetched_at=now_utc(),
        raw={k: (str(v) if v is not None else None) for k, v in row.items()},
    )


def _dt(v: Any):
    if v is None or v == "":
        return None
    if hasattr(v, "isoformat") and not isinstance(v, str):
        from ..timeutil import IST
        return v if getattr(v, "tzinfo", None) else v.replace(tzinfo=IST)
    return parse_portal_datetime(str(v))


def _amount(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v) if v > 0 else None
    return parse_inr_amount(str(v))


class ManualImportConnector(TenderSourceConnector):
    """Reads every CSV/XLSX/YAML file in ``config.path``."""

    supports_documents = True

    @property
    def directory(self) -> Path:
        if self.config.path is None:
            raise ValueError(f"source {self.source_id}: 'path' is required")
        return Path(self.config.path)

    def files(self) -> list[Path]:
        if not self.directory.exists():
            return []
        return sorted(p for p in self.directory.iterdir()
                      if p.is_file() and p.suffix.lower() in (".csv", ".xlsx", ".yaml", ".yml"))

    def discover(self, ctx: DiscoveryContext) -> Iterator[TenderRecord]:
        for f in self.files():
            ctx.pages_fetched += 1
            for row in _read_rows(f):
                try:
                    yield row_to_record(row, self.source_id, self.config.platform_label or "Manual import", f.parent, f.name)
                except (ValueError, TypeError) as exc:
                    ctx.malformed += 1
                    ctx.notes.append(f"{f.name}: skipped malformed row ({exc})")

    def fetch_listing(self, url: str | None, page_no: int) -> ListingPage:
        ctx = DiscoveryContext()
        return ListingPage(records=list(self.discover(ctx)), page_no=1)


class FixtureConnector(ManualImportConnector):
    """Offline connector for demos/tests: YAML/CSV records plus saved portal HTML
    (``*.cppp.html`` parsed with the CPPP parser, ``*.gepnic.html`` with GePNIC)."""

    def discover(self, ctx: DiscoveryContext) -> Iterator[TenderRecord]:
        from ..config import SourceConfig
        from .cppp import CPPPConnector
        from .gepnic import GePNICConnector

        yield from super().discover(ctx)
        for f in sorted(self.directory.glob("*.html")) if self.directory.exists() else []:
            ctx.pages_fetched += 1
            html = f.read_text(encoding="utf-8")
            if f.name.endswith(".cppp.html"):
                feed = "mmpdata" if "mmp" in f.name else ("gemdata" if "gem" in f.name else "cpppdata")
                conn = CPPPConnector(SourceConfig(id=self.source_id, type="cppp", feed=feed,
                                                  platform_label=self.config.platform_label))
                page = conn.parse_listing(html, f"https://eprocure.gov.in/cppp/latestactivetendersnew/{feed}")
            elif f.name.endswith(".gepnic.html"):
                conn = GePNICConnector(SourceConfig(id=self.source_id, type="gepnic",
                                                    base_url="https://eprocure.gov.in/eprocure/app",
                                                    platform_label=self.config.platform_label))
                page = conn.parse_listing(html, "https://eprocure.gov.in/eprocure/app")
            else:
                continue
            for rec in page.records:
                if ctx.since and rec.published_at and rec.published_at < ctx.since:
                    continue
                yield rec

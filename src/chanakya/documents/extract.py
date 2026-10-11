"""Stage 6 — document processing.

Native text first (pdfplumber / python-docx / openpyxl / HTML); OCR only for
pages with no text layer and only when enabled. Every piece of text keeps its
document name and page number so that downstream evidence can be cited.
Whole documents are never sent to an LLM: ``build_context`` assembles a
token-bounded excerpt from the prioritised sections.
"""

from __future__ import annotations

import hashlib
import logging
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from ..config import DocumentSettings
from ..security import SecurityError, ZipLimits, safe_extract_zip

log = logging.getLogger(__name__)

PRIORITY_SECTIONS = [
    "scope_of_work", "technical_requirements", "eligibility", "experience", "financial_qualification",
    "emd", "performance_security", "evaluation", "deliverables", "boq", "timeline", "key_dates", "contract_conditions",
]

SECTION_PATTERNS: dict[str, re.Pattern[str]] = {
    "scope_of_work": re.compile(r"\b(scope of (the )?work|scope of services|terms of reference|statement of work|project scope)\b", re.I),
    "technical_requirements": re.compile(r"\b(technical (specifications?|requirements?)|functional requirements?|system requirements?|specifications?)\b", re.I),
    "eligibility": re.compile(r"\b(eligibility criteria|pre[- ]?qualification( criteria)?|qualification criteria|minimum eligibility|bidder'?s? eligibility)\b", re.I),
    "experience": re.compile(r"\b(similar works?|past experience|work experience|experience criteria|completed projects?)\b", re.I),
    "financial_qualification": re.compile(r"\b(financial (criteria|capability|qualification)|annual turnover|average turnover|net worth)\b", re.I),
    "emd": re.compile(r"\b(earnest money( deposit)?|e\.?m\.?d\.?|bid security)\b", re.I),
    "performance_security": re.compile(r"\b(performance (security|bank guarantee|guarantee)|security deposit)\b", re.I),
    "evaluation": re.compile(r"\b(evaluation (criteria|methodology|process)|qcbs|technical evaluation|marking scheme)\b", re.I),
    "deliverables": re.compile(r"\b(deliverables|payment (terms|schedule|milestones?))\b", re.I),
    "boq": re.compile(r"\b(bill of quantit(y|ies)|b\.?o\.?q\.?|schedule of (quantities|rates|prices)|price (schedule|bid)|financial bid)\b", re.I),
    "timeline": re.compile(r"\b(completion period|project (timeline|schedule|duration)|contract period|period of contract|implementation schedule)\b", re.I),
    "key_dates": re.compile(r"\b(critical dates|key dates|important dates|bid submission (end )?date|last date (of|for) submission)\b", re.I),
    "contract_conditions": re.compile(r"\b(general conditions of contract|special conditions of contract|gcc|scc|liquidated damages|penalty clause)\b", re.I),
}

HEADING_LINE = re.compile(r"^\s*(?:(?:section|chapter|annexure|part)\s+[\w.-]+[:.)-]?\s*)?(?:\d+(?:\.\d+)*[.)]?\s+)?([A-Z][A-Za-z /&,'()-]{3,90})\s*:?\s*$")


@dataclass
class ExtractedPage:
    document: str
    page: int
    text: str
    method: str  # native | ocr | docx | xlsx | html | text


@dataclass
class Section:
    name: str
    document: str
    page_start: int
    page_end: int
    heading: str
    text: str

    def citation(self) -> str:
        pages = f"p.{self.page_start}" if self.page_start == self.page_end else f"pp.{self.page_start}-{self.page_end}"
        return f"{self.document} {pages} §{self.heading[:60]}"


@dataclass
class DocumentText:
    filename: str
    sha256: str
    mime: str
    pages: list[ExtractedPage] = field(default_factory=list)
    quality: float = 0.0
    needs_manual_review: bool = False
    method: str = "native"
    errors: list[str] = field(default_factory=list)
    children: list[DocumentText] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def all_pages(self) -> list[ExtractedPage]:
        out = list(self.pages)
        for c in self.children:
            out.extend(c.all_pages())
        return out


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def detect_type(path: Path) -> str:
    """Detect by magic bytes, not by (untrusted) extension."""
    with open(path, "rb") as fh:
        head = fh.read(8)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        import zipfile
        try:
            with zipfile.ZipFile(path) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc/.xls — not parsed, manual review
    sample = path.read_bytes()[:2048].lower()
    if b"<html" in sample or b"<!doctype html" in sample or b"<table" in sample:
        return "html"
    try:
        sample.decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "unknown"


def _quality(pages: list[ExtractedPage], min_chars: int) -> float:
    if not pages:
        return 0.0
    good = 0
    for p in pages:
        t = p.text.strip()
        if len(t) >= min_chars:
            alpha = sum(ch.isalpha() for ch in t) / max(1, len(t))
            good += 1 if alpha > 0.45 else 0.5
    return round(good / len(pages), 3)


def _ocr_page(pdf_path: Path, index: int, lang: str) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:  # pragma: no cover - optional dependency
        return ""
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        image = pdf[index].render(scale=2).to_pil()
        return pytesseract.image_to_string(image, lang=lang)
    finally:
        pdf.close()


def _extract_pdf(path: Path, name: str, settings: DocumentSettings, max_pages: int) -> tuple[list[ExtractedPage], list[str]]:
    import pdfplumber

    from .tables import TableState, page_text_with_tables

    pages: list[ExtractedPage] = []
    errors: list[str] = []
    tables = TableState(settings.max_table_rows_per_document) if settings.tables_enabled else None
    with pdfplumber.open(str(path)) as pdf:
        total = len(pdf.pages)
        if total > max_pages:
            errors.append(f"only first {max_pages} of {total} pages processed")
        for i, page in enumerate(pdf.pages[:max_pages]):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001 - malformed page content
                errors.append(f"page {i + 1}: {type(exc).__name__}")
                text = ""
            if tables is not None and text.strip():
                # Ruled tables (BOQs) become delimited rows; None means "no table here", keep the plain text.
                text = page_text_with_tables(page, i + 1, tables) or text
            method = "native"
            if len(text.strip()) < settings.min_chars_per_page and page.images and settings.ocr_enabled:
                ocr = _ocr_page(path, i, settings.ocr_languages)
                if len(ocr.strip()) > len(text.strip()):
                    text, method = ocr, "ocr"
            pages.append(ExtractedPage(name, i + 1, text, method))
    if tables is not None and tables.capped:
        errors.append(f"table rows capped at {settings.max_table_rows_per_document} per document: "
                      f"later table rows kept as plain text")
    return pages, errors


def _extract_docx(path: Path, name: str) -> list[ExtractedPage]:
    import docx

    d = docx.Document(str(path))
    blocks: list[str] = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            blocks.append(" | ".join(c.text.strip() for c in row.cells))
    # DOCX has no fixed pages; chunk into pseudo-pages of ~40 blocks for citations.
    chunks = [blocks[i:i + 40] for i in range(0, len(blocks), 40)] or [[]]
    return [ExtractedPage(name, i + 1, "\n".join(c), "docx") for i, c in enumerate(chunks)]


def _extract_xlsx(path: Path, name: str, max_rows: int = 2000) -> list[ExtractedPage]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets):
        lines = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= max_rows:
                lines.append(f"... truncated after {max_rows} rows")
                break
            vals = [str(v) for v in row if v not in (None, "")]
            if vals:
                lines.append(" | ".join(vals))
        pages.append(ExtractedPage(name, i + 1, f"[sheet {ws.title}]\n" + "\n".join(lines), "xlsx"))
    wb.close()
    return pages


def _extract_html(path: Path, name: str) -> list[ExtractedPage]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style", "nav", "header", "footer"]):
        tag.decompose()
    return [ExtractedPage(name, 1, soup.get_text("\n", strip=True), "html")]


def extract_document(path: Path, settings: DocumentSettings, *, display_name: str | None = None,
                     max_pages: int = 80, _depth: int = 0) -> DocumentText:
    name = display_name or path.name
    kind = detect_type(path)
    doc = DocumentText(filename=name, sha256=sha256_file(path), mime=kind)
    try:
        if kind == "pdf":
            doc.pages, doc.errors = _extract_pdf(path, name, settings, max_pages)
            doc.method = "ocr" if any(p.method == "ocr" for p in doc.pages) else "native"
        elif kind == "docx":
            doc.pages, doc.method = _extract_docx(path, name), "docx"
        elif kind == "xlsx":
            doc.pages, doc.method = _extract_xlsx(path, name), "xlsx"
        elif kind == "html":
            doc.pages, doc.method = _extract_html(path, name), "html"
        elif kind == "text":
            doc.pages, doc.method = [ExtractedPage(name, 1, path.read_text(encoding="utf-8", errors="replace"), "text")], "text"
        elif kind == "zip":
            if _depth >= 1:
                doc.errors.append("nested archive not expanded")
            else:
                tmp = Path(tempfile.mkdtemp(prefix="chanakya_zip_"))
                try:
                    limits = ZipLimits(settings.max_zip_members, settings.max_zip_total_mb * 1024 * 1024, settings.max_zip_ratio)
                    for member in safe_extract_zip(path, tmp, limits):
                        doc.children.append(extract_document(member, settings, display_name=f"{name}/{member.name}",
                                                             max_pages=max_pages, _depth=_depth + 1))
                finally:
                    shutil.rmtree(tmp, ignore_errors=True)
            doc.method = "zip"
        else:
            doc.errors.append(f"unsupported document type ({kind}); manual review required")
    except SecurityError as exc:
        doc.errors.append(f"rejected: {exc}")
    except Exception as exc:  # noqa: BLE001 - corrupted/malformed documents must not crash a run
        doc.errors.append(f"extraction failed: {type(exc).__name__}: {exc}")
    pages = doc.all_pages()
    doc.quality = _quality(pages, settings.min_chars_per_page)
    doc.needs_manual_review = doc.quality < 0.5 or (not pages) or any(e.startswith(("rejected", "extraction failed", "unsupported")) for e in doc.errors)
    return doc


def identify_sections(pages: list[ExtractedPage]) -> list[Section]:
    """Split pages into canonical sections using heading heuristics."""
    sections: list[Section] = []
    current: Section | None = None
    for page in pages:
        in_table = False  # rows inside a [TABLE n] ... [END TABLE n] block are data, never headings
        for line in page.text.splitlines():
            stripped = line.strip()
            if stripped.startswith("[TABLE "):
                in_table = True
            table_row, in_table = in_table, in_table and not stripped.startswith("[END TABLE")
            matched = None
            if not table_row and 4 <= len(stripped) <= 110 and (HEADING_LINE.match(stripped) or stripped.isupper() or stripped.endswith(":")):
                for name, pat in SECTION_PATTERNS.items():
                    if pat.search(stripped):
                        matched = name
                        break
            if matched:
                if current:
                    sections.append(current)
                current = Section(matched, page.document, page.page, page.page, stripped, "")
                continue
            if current is not None:
                if current.document != page.document:
                    sections.append(current)
                    current = None
                    continue
                current.text += line + "\n"
                current.page_end = page.page
    if current:
        sections.append(current)
    return [s for s in sections if s.text.strip()]


def build_context(sections: list[Section], max_chars: int, fallback_pages: list[ExtractedPage] | None = None) -> tuple[str, list[str]]:
    """Assemble a bounded, citation-tagged excerpt from prioritised sections.
    Returns (text, citations)."""
    order = {n: i for i, n in enumerate(PRIORITY_SECTIONS)}
    chosen = sorted(sections, key=lambda s: (order.get(s.name, 99), s.page_start))
    out, cites, used = [], [], 0
    per_section_cap = max(600, max_chars // 5)
    for s in chosen:
        if used >= max_chars:
            break
        body = re.sub(r"\n{2,}", "\n", s.text.strip())[: min(per_section_cap, max_chars - used)]
        tag = f"[{s.citation()}]"
        out.append(f"{tag}\n{body}")
        cites.append(s.citation())
        used += len(body) + len(tag) + 2
    if not out and fallback_pages:
        for p in fallback_pages:
            if used >= max_chars:
                break
            body = p.text.strip()[: max_chars - used]
            if body:
                out.append(f"[{p.document} p.{p.page}]\n{body}")
                cites.append(f"{p.document} p.{p.page}")
                used += len(body)
    return "\n\n".join(out), cites

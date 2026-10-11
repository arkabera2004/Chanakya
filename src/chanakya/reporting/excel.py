"""Excel workbook export (openpyxl, MIT licence).

Eight sheets: Executive Summary, All Opportunities, High Priority, Closing Soon,
Product-Wise, Eligibility Review, Change History, Scan Statistics.
All untrusted text is neutralised against formula injection; links are written
as hyperlinks only when they point to an allowlisted official domain.
"""

from __future__ import annotations

import hashlib
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..config import Settings
from ..db import AIUsageEvent, ChangeEvent, EligibilityFindingRow, ExportedReport, Opportunity, ScanRun, SourceScanRun
from ..models import TenderStatus
from ..pipeline.knowledge import KnowledgeBase
from ..security import is_safe_link, neutralize_cell
from ..timeutil import days_remaining, format_ist, from_storage, now_utc
from .insights import product_groups, relevant_book

MISSING = "Not stated"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PRIORITY_FILL = {
    "P1": PatternFill("solid", fgColor="C6EFCE"),
    "P2": PatternFill("solid", fgColor="FFEB9C"),
    "P3": PatternFill("solid", fgColor="DDEBF7"),
    "P4": PatternFill("solid", fgColor="EDEDED"),
}
MISSING_FONT = Font(italic=True, color="808080")
FLAG_FONT = Font(color="9C0006", bold=True)
INR_FORMAT = '"₹"#,##0'
DATE_FORMAT = "dd-mmm-yyyy hh:mm AM/PM"

ALL_COLUMNS = [
    ("Internal Opportunity ID", 16), ("Tender ID / Bid Number", 22), ("Tender Title", 60), ("Government Organization", 32),
    ("Department", 30), ("State / Region", 16), ("Procurement Platform", 28), ("Opportunity Type", 14),
    ("Publication Date (IST)", 20), ("Closing Date (IST)", 20), ("Days Remaining", 10), ("Estimated Value", 16),
    ("Currency", 9), ("Scope Summary", 50), ("Relevant Wayam Product", 28), ("Capability Match", 34),
    ("Relevance Score", 10), ("Priority", 9), ("Eligibility Status", 26), ("Eligibility Requirements", 50),
    ("Turnover Requirement", 26), ("Past Experience Requirement", 34), ("Certification Requirements", 26), ("EMD", 14),
    ("Hardware Requirement", 14), ("Partner Requirement", 18), ("Technical Risks", 40), ("Business Recommendation", 46),
    ("Suggested Next Action", 50), ("Official Tender URL", 40), ("Tender Document URL", 34), ("Evidence Reference", 50),
    ("Confidence", 10), ("Last Verified (IST)", 20), ("Current Tender Status", 14),
    # extra transparency columns (beyond the required 35)
    ("Flags", 40), ("Score Breakdown", 60), ("Reference No.", 26), ("Data Completeness", 10), ("AI Analysis", 18),
    ("Review Status", 12),
]


def _ist_naive(dt):
    """Excel has no time zones: write IST wall-clock time as a naive datetime."""
    return from_storage(dt).replace(tzinfo=None) if dt else None


def _write_header(ws, columns: list[tuple[str, int]]) -> None:
    for i, (name, width) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _put(ws, row: int, col: int, value, *, link: str | None = None, fmt: str | None = None, wrap: bool = False):
    if value is None or value == "" or value == []:
        c = ws.cell(row=row, column=col, value=MISSING)
        c.font = MISSING_FONT
        return c
    if isinstance(value, list):
        value = "; ".join(str(v) for v in value)
    c = ws.cell(row=row, column=col, value=neutralize_cell(value))
    if isinstance(c.value, str):
        c.data_type = "s"
    if link:
        c.hyperlink = link
        c.font = Font(color="0563C1", underline="single")
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def _finish(ws, ncols: int, nrows: int) -> None:
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(nrows, 1)}"


def _findings(session: Session, opp_ids: list[str]) -> dict[str, list[EligibilityFindingRow]]:
    out: dict[str, list[EligibilityFindingRow]] = {i: [] for i in opp_ids}
    if not opp_ids:
        return out
    for row in session.scalars(select(EligibilityFindingRow).where(EligibilityFindingRow.opportunity_id.in_(opp_ids))):
        out.setdefault(row.opportunity_id, []).append(row)
    return out


def _req_text(rows, criterion: str) -> str | None:
    items = [f"{r.requirement_text[:160]} [{r.assessment}{'' if r.explicit else ', inferred'}]" for r in rows if r.criterion == criterion]
    return "; ".join(items) if items else None


def _breakdown_text(b: dict) -> str:
    comps = b.get("components", {})
    return "; ".join(f"{k} {v.get('points', 0):.1f}/{v.get('weight', 0):.0f} ({v.get('reason', '')})" for k, v in comps.items())


def select_opportunities(session: Session, include_rejected: bool = False, active_only: bool = False) -> list[Opportunity]:
    q = select(Opportunity).where(Opportunity.relevance_score.is_not(None))
    if not include_rejected:
        q = q.where(Opportunity.filter_decision != "rejected")
    if active_only:
        q = q.where(Opportunity.status == TenderStatus.ACTIVE.value)
    rows = list(session.scalars(q).all())
    rows.sort(key=lambda o: (o.priority or "P9", -(o.relevance_score or 0), o.closing_at or now_utc() + timedelta(days=9999)))
    return rows


def export_workbook(session: Session, settings: Settings, kb: KnowledgeBase, run: ScanRun | None,
                    out_dir: Path | None = None) -> tuple[Path, int]:
    out_dir = out_dir or settings.app.reports_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    allowed = settings.http.allowed_domains
    opps = select_opportunities(session)
    findings = _findings(session, [o.id for o in opps])
    now = now_utc()
    wb = Workbook()

    # ---------------------------------------------------------- 1 summary
    ws = wb.active
    ws.title = "Executive Summary"
    stats = (run.stats if run else {}) or {}
    active = [o for o in opps if o.status == TenderStatus.ACTIVE.value]
    relevant = relevant_book(active)
    soon = [o for o in relevant if o.closing_at and 0 <= (days_remaining(o.closing_at, now) or -1) <= settings.scoring.closing_soon_days]
    review = [o for o in active if o.priority in ("P1", "P2") and o.eligibility_status != "Verified eligible"]
    not_active = session.scalar(select(func.count()).select_from(Opportunity).where(Opportunity.status != "active")) or 0
    rejected = session.scalar(select(func.count()).select_from(Opportunity).where(Opportunity.filter_decision == "rejected")) or 0
    src_stats = stats.get("sources", {})
    usage_q = select(func.count(), func.coalesce(func.sum(AIUsageEvent.input_tokens), 0),
                     func.coalesce(func.sum(AIUsageEvent.output_tokens), 0),
                     func.coalesce(func.sum(AIUsageEvent.est_cost_usd), 0.0))
    if run:
        usage_q = usage_q.where(AIUsageEvent.run_id == run.id, AIUsageEvent.cache_hit.is_(False),
                                AIUsageEvent.result.notin_(["budget_blocked"]))
    calls, tin, tout, cost = session.execute(usage_q).one()
    rows = [
        ("Report generated (IST)", format_ist(now)),
        ("Run ID", run.id if run else "n/a (export without run)"),
        ("Run status", run.status if run else "n/a"),
        ("Sources checked", f"{len(src_stats)} ({sum(1 for v in src_stats.values() if v.get('status') == 'ok')} ok, "
                            f"{sum(1 for v in src_stats.values() if v.get('status') in ('failed', 'blocked'))} failed/blocked, "
                            f"{sum(1 for v in src_stats.values() if v.get('status') == 'skipped')} skipped)"),
        ("New opportunities discovered (this run)", stats.get("new", 0)),
        ("Updated opportunities (this run)", stats.get("updated", 0)),
        ("Total relevant active opportunities (P1 to P3)", len(relevant)),
        ("High-priority (P1)", sum(1 for o in active if o.priority == "P1")),
        ("Worth investigating (P2)", sum(1 for o in active if o.priority == "P2")),
        (f"Closing within {settings.scoring.closing_soon_days} days", len(soon)),
        ("Expired / withdrawn / closed (all time)", not_active),
        ("Requiring eligibility review (P1/P2 not verified)", len(review)),
        ("Filtered out as unrelated (deterministic, not listed)", rejected),
        ("AI mode", stats.get("llm_mode", settings.llm.mode)),
        ("AI calls (billable, this run)", calls),
        ("AI input / output tokens", f"{tin:,} / {tout:,}"),
        ("Estimated AI cost (USD, this run)", round(float(cost), 4)),
        ("AI analyses skipped by budget", stats.get("llm_skipped_budget", 0)),
        ("Deep analyses deferred (no documents)", stats.get("deep_deferred", 0)),
        ("Coverage notes", "\n".join(stats.get("coverage_notes", [])) or "None"),
        ("Important", "Scores are decision support. Eligibility is never assumed: 'Unknown' means Wayam's credentials "
                      "were not available to verify. Verify every deadline and requirement on the official portal."),
    ]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 48, 100
    ws.cell(row=1, column=1, value="CHANAKYA — Wayam AI Government Opportunities").font = Font(bold=True, size=14)
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = _put(ws, i, 2, v, wrap=True)
        if k == "Important":
            c.font = FLAG_FONT

    # ---------------------------------------------------------- 2 all opportunities
    ws = wb.create_sheet("All Opportunities")
    _write_header(ws, ALL_COLUMNS)
    for r, o in enumerate(opps, start=2):
        intel = o.intelligence or {}
        f = findings.get(o.id, [])
        days = days_remaining(o.closing_at, now)
        caps = "; ".join(f"{c['id']} ({c['score']:.2f})" for c in (o.top_capabilities or []))
        elig_reqs = "; ".join(f"{x.criterion}: {x.requirement_text[:120]} [{x.assessment}]" for x in f
                              if x.criterion not in ("emd",))
        emd = o.emd_inr or next((x.value for x in f if x.criterion == "emd" and x.value), None)
        partner = "Yes" if o.eligibility_status == "Partner required" else ("Possibly" if intel.get("fit") == "partner_needed" else None)
        hw = {True: "Yes", False: "No"}.get(intel.get("hardware_involved"))
        doc_note = None
        refs = o.document_refs or []
        direct = next((d.get("url") for d in refs if d.get("url")), None)
        if direct and is_safe_link(direct, allowed):
            doc_note = direct
        elif any(d.get("access") == "captcha_required" for d in refs):
            doc_note = "Portal download (CAPTCHA) — fetch manually"
        values = [
            o.id, o.source_tender_id, o.title, o.organisation, o.department, o.state, o.platform, o.opportunity_type,
            _ist_naive(o.published_at), _ist_naive(o.closing_at), round(days, 1) if days is not None else None,
            o.estimated_value_inr, "INR" if o.estimated_value_inr else None, intel.get("summary"),
            intel.get("relevant_solution"), caps, o.relevance_score, o.priority, o.eligibility_status, elig_reqs or None,
            _req_text(f, "turnover"), _req_text(f, "experience"), _req_text(f, "certification"), emd, hw, partner,
            intel.get("technical_risks"), intel.get("why_match"), intel.get("recommended_next_action"),
            o.official_url, doc_note, intel.get("evidence"), intel.get("confidence"), _ist_naive(o.last_verified_at),
            o.status, intel.get("flags"), _breakdown_text(o.score_breakdown or {}), o.ref_no, o.data_completeness,
            f"{o.analysis_state} ({intel.get('confidence_basis', '')})", o.review_status,
        ]
        for c, v in enumerate(values, start=1):
            header = ALL_COLUMNS[c - 1][0]
            link = v if header in ("Official Tender URL", "Tender Document URL") and isinstance(v, str) and is_safe_link(v, allowed) else None
            fmt = DATE_FORMAT if "Date" in header or "Verified" in header else (INR_FORMAT if header in ("Estimated Value", "EMD") else None)
            cell = _put(ws, r, c, v, link=link, fmt=fmt, wrap=header in ("Tender Title", "Scope Summary", "Business Recommendation",
                                                                          "Suggested Next Action", "Evidence Reference", "Flags"))
            if header == "Priority" and o.priority in PRIORITY_FILL:
                cell.fill = PRIORITY_FILL[o.priority]
                cell.font = Font(bold=True)
            if header == "Flags" and v:
                cell.font = FLAG_FONT
    _finish(ws, len(ALL_COLUMNS), len(opps) + 1)

    # ---------------------------------------------------------- 3 high priority / 4 closing soon
    brief_cols = [("Internal Opportunity ID", 16), ("Priority", 9), ("Relevance Score", 10), ("Tender Title", 60),
                  ("Government Organization", 30), ("Closing Date (IST)", 20), ("Days Remaining", 10),
                  ("Relevant Wayam Product", 28), ("Eligibility Status", 26), ("Flags", 40),
                  ("Suggested Next Action", 50), ("Tender ID / Bid Number", 22), ("Official Tender URL", 40)]

    def brief(ws_name: str, items: list[Opportunity]) -> None:
        w = wb.create_sheet(ws_name)
        _write_header(w, brief_cols)
        for r, o in enumerate(items, start=2):
            intel = o.intelligence or {}
            d = days_remaining(o.closing_at, now)
            vals = [o.id, o.priority, o.relevance_score, o.title, o.organisation, _ist_naive(o.closing_at),
                    round(d, 1) if d is not None else None, intel.get("relevant_solution"), o.eligibility_status,
                    intel.get("flags"), intel.get("recommended_next_action"), o.source_tender_id, o.official_url]
            for c, v in enumerate(vals, start=1):
                h = brief_cols[c - 1][0]
                cell = _put(w, r, c, v, link=v if h == "Official Tender URL" and is_safe_link(v, allowed) else None,
                            fmt=DATE_FORMAT if "Date" in h else None, wrap=h in ("Tender Title", "Suggested Next Action", "Flags"))
                if h == "Priority" and o.priority in PRIORITY_FILL:
                    cell.fill = PRIORITY_FILL[o.priority]
                if h == "Flags" and v:
                    cell.font = FLAG_FONT
        _finish(w, len(brief_cols), len(items) + 1)

    brief("High Priority", [o for o in active if o.priority in ("P1", "P2")])
    brief("Closing Soon", sorted(soon, key=lambda o: o.closing_at))

    # ---------------------------------------------------------- 5 product-wise
    ws = wb.create_sheet("Product-Wise")
    pcols = [("Wayam Offering", 34), ("Capability ID", 18), ("Offering Status", 16), ("Internal Opportunity ID", 16),
             ("Priority", 9), ("Relevance Score", 10), ("Capability Score", 10), ("Tender Title", 60),
             ("Government Organization", 30), ("Closing Date (IST)", 20), ("Eligibility Status", 26)]
    _write_header(ws, pcols)
    prow = 2
    for cap, items in product_groups(relevant, kb):  # the desk Products page reads the same grouping
        for o, c in items:
            vals = [cap.offering, cap.id, cap.offering_status, o.id, o.priority, o.relevance_score, c["score"], o.title,
                    o.organisation, _ist_naive(o.closing_at), o.eligibility_status]
            for ci, v in enumerate(vals, start=1):
                cell = _put(ws, prow, ci, v, fmt=DATE_FORMAT if ci == 10 else None, wrap=ci == 8)
                if ci == 5 and o.priority in PRIORITY_FILL:
                    cell.fill = PRIORITY_FILL[o.priority]
            prow += 1
    _finish(ws, len(pcols), prow - 1)

    # ---------------------------------------------------------- 6 eligibility review
    ws = wb.create_sheet("Eligibility Review")
    ecols = [("Internal Opportunity ID", 16), ("Priority", 9), ("Tender Title", 50), ("Eligibility Status", 26),
             ("Criterion", 18), ("Requirement (as stated)", 70), ("Explicit / Inferred", 12), ("Assessment", 16),
             ("Evidence", 50), ("Eligibility Blockers", 50)]
    _write_header(ws, ecols)
    er = 2
    for o in [o for o in active if o.priority in ("P1", "P2", "P3")
              and o.eligibility_status != "Verified eligible"]:
        rows_f = findings.get(o.id) or [None]
        for fnd in rows_f:
            ev = None
            if fnd and fnd.evidence:
                e = fnd.evidence
                ev = " ".join(str(x) for x in (e.get("source"), e.get("document"), f"p.{e['page']}" if e.get("page") else None,
                                               e.get("section"), f"“{e['quote']}”" if e.get("quote") else None) if x)
            vals = [o.id, o.priority, o.title, o.eligibility_status,
                    fnd.criterion if fnd else "none extracted",
                    fnd.requirement_text if fnd else "No eligibility text available — tender documents not reviewed",
                    ("Explicit" if fnd.explicit else "Inferred") if fnd else None, fnd.assessment if fnd else "unknown",
                    ev, (o.intelligence or {}).get("eligibility_blockers")]
            for ci, v in enumerate(vals, start=1):
                _put(ws, er, ci, v, wrap=ci in (3, 6, 9, 10))
            er += 1
    _finish(ws, len(ecols), er - 1)

    # ---------------------------------------------------------- 7 change history
    ws = wb.create_sheet("Change History")
    ccols = [("Detected (IST)", 20), ("Internal Opportunity ID", 16), ("Tender Title", 50), ("Change Type", 18),
             ("Field", 16), ("Old Value", 30), ("New Value", 30), ("Material", 9), ("Run ID", 26)]
    _write_header(ws, ccols)
    since = now - timedelta(days=30)
    q = (select(ChangeEvent, Opportunity).join(Opportunity, Opportunity.id == ChangeEvent.opportunity_id)
         .where(ChangeEvent.detected_at >= since, Opportunity.filter_decision != "rejected")
         .order_by(ChangeEvent.detected_at.desc()).limit(5000))
    cr = 2
    for ev, o in session.execute(q):
        vals = [_ist_naive(ev.detected_at), o.id, o.title, ev.change_type, ev.field, ev.old_value, ev.new_value,
                "Yes" if ev.material else "No", ev.run_id]
        for ci, v in enumerate(vals, start=1):
            _put(ws, cr, ci, v, fmt=DATE_FORMAT if ci == 1 else None, wrap=ci == 3)
        cr += 1
    _finish(ws, len(ccols), cr - 1)

    # ---------------------------------------------------------- 8 scan statistics
    ws = wb.create_sheet("Scan Statistics")
    scols = [("Source", 22), ("Status", 10), ("Started (IST)", 20), ("Duration (s)", 10), ("Pages", 8),
             ("Listings Seen", 10), ("New", 8), ("Updated", 8), ("Unchanged", 10), ("Duplicates", 10),
             ("Duplicate Rate", 10), ("Malformed", 10), ("Error / Coverage Note", 70)]
    _write_header(ws, scols)
    sr = 2
    if run:
        for s in session.scalars(select(SourceScanRun).where(SourceScanRun.run_id == run.id)):
            dup_rate = round(s.duplicate_count / s.listings_seen, 3) if s.listings_seen else None
            vals = [s.source_id, s.status, _ist_naive(s.started_at), s.duration_s, s.pages_fetched, s.listings_seen,
                    s.new_count, s.updated_count, s.unchanged_count, s.duplicate_count, dup_rate, s.malformed_count,
                    " | ".join(x for x in (s.error, s.coverage_note) if x) or "None"]
            for ci, v in enumerate(vals, start=1):
                _put(ws, sr, ci, v, fmt=DATE_FORMAT if ci == 3 else None, wrap=ci == 13)
            sr += 1
    sr += 1
    ws.cell(row=sr, column=1, value="AI usage by model and purpose (this run)").font = Font(bold=True)
    sr += 1
    ucols = ["Model", "Purpose", "Result", "Calls", "Input tokens", "Output tokens", "Cached tokens", "Est. cost (USD)"]
    for ci, h in enumerate(ucols, start=1):
        c = ws.cell(row=sr, column=ci, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    sr += 1
    if run:
        uq = (select(AIUsageEvent.model, AIUsageEvent.purpose, AIUsageEvent.result, func.count(),
                     func.sum(AIUsageEvent.input_tokens), func.sum(AIUsageEvent.output_tokens),
                     func.sum(AIUsageEvent.cached_tokens), func.sum(AIUsageEvent.est_cost_usd))
              .where(AIUsageEvent.run_id == run.id)
              .group_by(AIUsageEvent.model, AIUsageEvent.purpose, AIUsageEvent.result))
        for row in session.execute(uq):
            for ci, v in enumerate(row, start=1):
                _put(ws, sr, ci, round(v, 6) if isinstance(v, float) else v)
            sr += 1

    stamp = now.replace(tzinfo=None) + timedelta(hours=5, minutes=30)
    name = f"chanakya_opportunities_{stamp:%Y-%m-%d_%H%M}" + (f"_{run.id[-4:]}" if run else "") + ".xlsx"
    path = out_dir / name
    wb.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    rep = ExportedReport(run_id=run.id if run else None, path=str(path), sha256=digest, rows=len(opps))
    session.add(rep)
    session.commit()
    _apply_retention(out_dir, settings.app.report_retention)
    return path, rep.id


def _apply_retention(out_dir: Path, keep: int) -> None:
    files = sorted((p for p in out_dir.glob("chanakya_opportunities_*.xlsx")), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in files[keep:]:
        old.unlink(missing_ok=True)


def latest_report(session: Session) -> ExportedReport | None:
    return session.scalar(select(ExportedReport).order_by(ExportedReport.id.desc()))


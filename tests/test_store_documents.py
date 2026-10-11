"""Deduplication, change tracking and document processing."""

from __future__ import annotations

import io
import zipfile

import pytest
from sqlalchemy import select

from chanakya.config import DocumentSettings
from chanakya.db import ChangeEvent, Opportunity, OpportunityVersion
from chanakya.documents import build_context, extract_document, identify_sections
from chanakya.models import TenderRecord, TenderStatus
from chanakya.pipeline.store import dedup_key, mark_expired, opportunity_id, upsert_record
from chanakya.timeutil import parse_portal_datetime

from .conftest import ist


def rec(**kw) -> TenderRecord:
    d = dict(source_id="cppp_central", platform="CPPP", source_record_key="2026_ABC_1_1", source_tender_id="2026_ABC_1_1",
             ref_no="ABC/1", title="AI video analytics for city cameras", organisation="Synthetic Org",
             published_at=parse_portal_datetime(ist(-1)), closing_at=parse_portal_datetime(ist(20)))
    d.update(kw)
    return TenderRecord(**d)


# ---------------------------------------------------------------- dedup keys
def test_dedup_key_rules():
    assert dedup_key(rec()) == "tid:2026_ABC_1_1"
    assert dedup_key(rec(source_tender_id="171861", organisation="CPWD")).startswith("sid:cpwd:")
    assert dedup_key(rec(source_tender_id=None)).startswith("ref:")
    assert dedup_key(rec(source_tender_id=None, ref_no=None)).startswith("h:")
    assert opportunity_id("tid:X") == opportunity_id("tid:X") and opportunity_id("tid:X").startswith("CHK-")


def test_new_unchanged_updated_and_deadline_change(session):
    r1 = upsert_record(session, rec(), "run1")
    session.commit()
    assert r1.outcome == "new"
    assert upsert_record(session, rec(), "run2").outcome == "unchanged"
    r3 = upsert_record(session, rec(closing_at=parse_portal_datetime(ist(30)), has_corrigendum=True,
                                    corrigenda=["Corrigendum 1: date extension"]), "run3")
    session.commit()
    assert r3.outcome == "updated" and r3.opportunity.version == 2
    types = {c.change_type for c in r3.changes}
    assert {"deadline_changed", "corrigendum_added"} <= types
    assert session.scalar(select(OpportunityVersion).where(OpportunityVersion.version == 2)) is not None


def test_cross_portal_duplicate_is_merged(session):
    upsert_record(session, rec(), "run1")
    session.commit()
    other = rec(source_id="manual", source_record_key="m-1", source_tender_id=None, platform="Manual",
                description="Scope details from the PDF")
    res = upsert_record(session, other, "run1")
    session.commit()
    assert res.outcome == "duplicate"
    opp = session.scalar(select(Opportunity))
    assert opp.also_seen_on[0]["source_id"] == "manual" and opp.description == "Scope details from the PDF"
    assert session.scalar(select(Opportunity).where(Opportunity.source_id == "manual")) is None


def test_fuzzy_title_duplicate_same_closing(session):
    upsert_record(session, rec(), "run1")
    session.commit()
    near = rec(source_id="other", source_record_key="x", source_tender_id=None, ref_no=None,
               title="AI video analytics for city cameras.")
    assert upsert_record(session, near, "run1").outcome == "duplicate"


def test_retender_detected(session):
    upsert_record(session, rec(source_tender_id="2026_ABC_77_1", source_record_key="2026_ABC_77_1"), "r1")
    session.commit()
    res = upsert_record(session, rec(source_tender_id="2026_ABC_77_2", source_record_key="2026_ABC_77_2", ref_no="ABC/77-R"), "r2")
    session.commit()
    assert any(c.change_type == "retendered" for c in res.changes)
    old = session.get(Opportunity, opportunity_id("tid:2026_ABC_77_1"))
    assert old.status == TenderStatus.RETENDERED.value


def test_mark_expired(session):
    upsert_record(session, rec(closing_at=parse_portal_datetime(ist(-1))), "r1")
    session.commit()
    assert mark_expired(session, "r2") == 1
    session.commit()
    assert session.scalar(select(Opportunity)).status == "expired"
    assert session.scalar(select(ChangeEvent).where(ChangeEvent.change_type == "status_changed")) is not None


# ---------------------------------------------------------------- documents
def make_pdf(path, pages: list[str]):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(50, y, line)
            y -= 16
        c.showPage()
    c.save()


RFP_PAGES = [
    "NOTICE INVITING TENDER\nSynthetic Smart City Ltd\nProject: AI Video Analytics",
    "1. Scope of Work\nDeploy video analytics on 450 existing cameras for crowd density and intrusion.\nIntegrate alerts with ICCC.",
    "2. Eligibility Criteria\nAverage annual turnover of Rs. 10 crore in the last three financial years.\n"
    "Bidder must be ISO 27001 certified.\nConsortium is not allowed.",
    "3. Earnest Money Deposit\nEMD of Rs. 9,00,000 through online payment.",
]


def test_text_pdf_extraction_with_sections_and_citations(tmp_path):
    p = tmp_path / "rfp.pdf"
    make_pdf(p, RFP_PAGES)
    doc = extract_document(p, DocumentSettings())
    assert doc.mime == "pdf" and doc.page_count == 4 and not doc.needs_manual_review
    secs = identify_sections(doc.all_pages())
    names = {s.name: s for s in secs}
    assert names["eligibility"].page_start == 3 and "turnover" in names["eligibility"].text.lower()
    assert names["scope_of_work"].page_start == 2
    ctx, cites = build_context(secs, max_chars=400)
    assert len(ctx) <= 600 and any("p.2" in c for c in cites) and cites[0].startswith("rfp.pdf")


def test_page_limit_and_quality(tmp_path):
    p = tmp_path / "long.pdf"
    make_pdf(p, [f"Page {i} technical specifications of the system and services required." for i in range(12)])
    doc = extract_document(p, DocumentSettings(), max_pages=5)
    assert doc.page_count == 5 and any("only first 5" in e for e in doc.errors)


def test_scanned_pdf_flagged_for_manual_review(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    p = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.rect(50, 50, 400, 600, fill=1)  # no text layer
    c.showPage()
    c.save()
    doc = extract_document(p, DocumentSettings(ocr_enabled=False))
    assert doc.needs_manual_review and doc.quality == 0


def test_corrupted_documents_do_not_crash(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.7\n this is not really a pdf")
    doc = extract_document(p, DocumentSettings())
    assert doc.needs_manual_review and doc.errors
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    assert "unsupported" in extract_document(legacy, DocumentSettings()).errors[0]


def test_docx_xlsx_html_and_zip(tmp_path):
    import docx
    from openpyxl import Workbook
    d = docx.Document()
    d.add_paragraph("Eligibility Criteria")
    d.add_paragraph("Average annual turnover of Rs. 50 lakh.")
    d.save(tmp_path / "annex.docx")
    wb = Workbook()
    wb.active.append(["Item", "Qty"])
    wb.active.append(["Video analytics licence", 450])
    wb.save(tmp_path / "boq.xlsx")
    (tmp_path / "notice.html").write_text("<html><body><nav>menu</nav><p>Scope of Work: chatbot</p></body></html>")
    assert "turnover" in extract_document(tmp_path / "annex.docx", DocumentSettings()).pages[0].text
    assert "Video analytics licence" in extract_document(tmp_path / "boq.xlsx", DocumentSettings()).pages[0].text
    html_doc = extract_document(tmp_path / "notice.html", DocumentSettings())
    assert "chatbot" in html_doc.pages[0].text and "menu" not in html_doc.pages[0].text
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.write(tmp_path / "annex.docx", "tender/annex.docx")
        zf.write(tmp_path / "boq.xlsx", "tender/boq.xlsx")
    (tmp_path / "bundle.zip").write_bytes(buf.getvalue())
    z = extract_document(tmp_path / "bundle.zip", DocumentSettings())
    assert len(z.children) == 2 and len(z.all_pages()) >= 2


def test_malicious_zip_rejected_without_crash(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../../escape.txt", "x")
    (tmp_path / "evil.zip").write_bytes(buf.getvalue())
    doc = extract_document(tmp_path / "evil.zip", DocumentSettings())
    assert doc.needs_manual_review and any("rejected" in e for e in doc.errors)
    assert not (tmp_path.parent / "escape.txt").exists()


@pytest.mark.parametrize("magic,kind", [(b"%PDF-1.4", "pdf"), (b"<html><table>", "html"), (b"plain text", "text")])
def test_detect_type_by_magic(tmp_path, magic, kind):
    from chanakya.documents.extract import detect_type
    p = tmp_path / "file.bin"   # extension deliberately misleading
    p.write_bytes(magic + b" content")
    assert detect_type(p) == kind

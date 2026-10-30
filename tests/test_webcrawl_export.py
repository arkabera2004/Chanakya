"""CSV/XLSX export of extracted WebsiteTenders. Exports must preserve the source URL and
evidence excerpt for every field — the brief's explicit requirement.
"""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from chanakya.webcrawl.export import export_csv, export_xlsx

TENDERS = [
    {
        "source_url": "https://example.gov.in/tenders/1",
        "fields": {
            "tender_id": {"value": "NHAI/2026/1", "excerpt": "Tender No: NHAI/2026/1", "method": "rule_based"},
            "state": {"value": "Tamil Nadu", "excerpt": "Tamil Nadu", "method": "rule_based"},
        },
        "matches": [{"capability_id": "CAP-AI-01", "score": 0.72, "strong": ["video analytics"], "medium": []}],
    },
]


def test_export_csv_preserves_source_url_and_evidence():
    text = export_csv(TENDERS)
    rows = list(csv.DictReader(io.StringIO(text)))
    assert len(rows) == 1
    assert rows[0]["source_url"] == "https://example.gov.in/tenders/1"
    assert rows[0]["tender_id"] == "NHAI/2026/1"
    assert rows[0]["tender_id_excerpt"] == "Tender No: NHAI/2026/1"
    assert rows[0]["top_capability"] == "CAP-AI-01"
    assert rows[0]["top_capability_score"] == "0.72"


def test_export_xlsx_preserves_source_url_and_evidence():
    data = export_xlsx(TENDERS)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    by_col = dict(zip(header, row, strict=True))
    assert by_col["source_url"] == "https://example.gov.in/tenders/1"
    assert by_col["tender_id"] == "NHAI/2026/1"
    assert by_col["tender_id_excerpt"] == "Tender No: NHAI/2026/1"


def test_export_with_no_tenders_still_has_headers():
    text = export_csv([])
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["source_url", "top_capability", "top_capability_score"]

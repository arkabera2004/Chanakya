"""Excel, email and end-to-end pipeline runs (fixture source + simulated LLM)."""

from __future__ import annotations

import email
import smtplib
from pathlib import Path

import httpx
import pytest
import yaml
from openpyxl import load_workbook
from sqlalchemy import func, select
from typer.testing import CliRunner

from chanakya.app import run_pipeline
from chanakya.cli import app as cli_app
from chanakya.config import HttpSettings, SourceConfig
from chanakya.db import AIUsageEvent, ChangeEvent, EmailDelivery, ExportedReport, Opportunity, ScanRun
from chanakya.http import PoliteClient
from chanakya.llm import FakeProvider
from chanakya.reporting.email import SMTPEmailProvider, build_digest, send_digest

from .conftest import ist, sample_records, write_records

REQUIRED_COLUMNS = [
    "Internal Opportunity ID", "Tender ID / Bid Number", "Tender Title", "Government Organization", "Department",
    "State / Region", "Procurement Platform", "Opportunity Type", "Publication Date (IST)", "Closing Date (IST)",
    "Days Remaining", "Estimated Value", "Currency", "Scope Summary", "Relevant Wayam Product", "Capability Match",
    "Relevance Score", "Priority", "Eligibility Status", "Eligibility Requirements", "Turnover Requirement",
    "Past Experience Requirement", "Certification Requirements", "EMD", "Hardware Requirement", "Partner Requirement",
    "Technical Risks", "Business Recommendation", "Suggested Next Action", "Official Tender URL", "Tender Document URL",
    "Evidence Reference", "Confidence", "Last Verified (IST)", "Current Tender Status",
]


def run(settings, session, **kw):
    kw.setdefault("provider", FakeProvider())
    return run_pipeline(settings, session, **kw)


def llm_calls(session, run_id) -> int:
    return session.scalar(select(func.count()).select_from(AIUsageEvent).where(
        AIUsageEvent.run_id == run_id, AIUsageEvent.cache_hit.is_(False), AIUsageEvent.result != "budget_blocked")) or 0


# ---------------------------------------------------------------- full pipeline
def test_full_run_then_incremental_runs(settings, session):
    src = settings.sources[0].path
    write_records(src, sample_records())
    r1 = run(settings, session, email=True)
    assert r1.run.status == "ok" and r1.exit_code == 0
    s = r1.stats
    assert s.new == 6 and s.assessed == 6 and s.expired_marked == 1
    assert s.rejected >= 1 and s.rejected + s.low >= 2       # catering rejected; road resurfacing kept LOW — both zero LLM cost

    opps = {o.source_tender_id: o for o in session.scalars(select(Opportunity))}
    va = opps["2026_SYNSC_000101_1"]
    assert va.priority in ("P1", "P2") and va.top_capabilities[0]["id"] == "CAP-CV-VA"
    assert va.eligibility_status != "Verified eligible"       # company profile is empty → never assumed
    assert va.intelligence["evidence"] and "Fixture portal" in va.intelligence["evidence"][0] or va.evidence
    assert set(va.score_breakdown["components"]) == {"capability", "eligibility", "technical", "deadline", "commercial", "strategic"}
    assert opps["2026_SYNCAT_000106_1"].filter_decision == "rejected" and opps["2026_SYNCAT_000106_1"].priority == "P4"
    assert opps["2026_SYNOLD_000105_1"].status == "expired" and opps["2026_SYNOLD_000105_1"].priority == "P4"
    assert any("SIMULATED" in f for f in va.intelligence["flags"]) or va.analysis_state in ("not_needed",)

    # rejected/expired records never reached the LLM
    called_for = {e.opportunity_id for e in session.scalars(select(AIUsageEvent).where(AIUsageEvent.run_id == r1.run.id))}
    assert opps["2026_SYNCAT_000106_1"].id not in called_for and opps["2026_SYNOLD_000105_1"].id not in called_for
    assert opps["2026_SYNPWD_000102_1"].id not in called_for

    # Excel + email produced and recorded
    assert r1.report_path and r1.report_path.exists()
    assert sorted(st for _, st in r1.email_status) == ["sent", "sent"]
    emls = list(settings.email.outbox_dir.glob("*.eml"))
    assert len(emls) == 2
    msg = email.message_from_bytes(emls[0].read_bytes())
    assert msg["Subject"].startswith("CHANAKYA | Wayam AI Government Opportunities |")
    assert any(part.get_filename() == r1.report_path.name for part in msg.walk())

    # run 2: nothing changed → no reprocessing, no LLM calls, no duplicate emails for the same report
    r2 = run(settings, session)
    assert (r2.stats.new, r2.stats.updated, r2.stats.unchanged) == (0, 0, 6)
    assert r2.stats.assessed == 0 and llm_calls(session, r2.run.id) == 0

    # run 3: a corrigendum extends one deadline → change detected and only that record re-assessed
    recs = sample_records()
    recs[0]["closing_at"] = ist(40)
    write_records(src, recs)
    r3 = run(settings, session)
    assert r3.stats.updated == 1 and r3.stats.assessed == 1
    ev = session.scalar(select(ChangeEvent).where(ChangeEvent.change_type == "deadline_changed"))
    assert ev is not None and ev.run_id == r3.run.id
    wb = load_workbook(r3.report_path)
    changes = [row[3] for row in wb["Change History"].iter_rows(min_row=2, values_only=True)]
    assert "deadline_changed" in changes


def test_deferred_analysis_resumes_when_documents_are_imported(settings, session):
    from chanakya.db import EligibilityFindingRow, OpportunityDocument
    from .test_store_documents import make_pdf
    write_records(settings.sources[0].path, sample_records())
    run(settings, session)
    voice = session.scalar(select(Opportunity).where(Opportunity.source_tender_id == "2026_SYNHLT_000103_1"))
    assert voice.analysis_state == "deferred_no_documents"
    folder = settings.app.inbox_dir / "documents" / "2026_SYNHLT_000103_1"
    folder.mkdir(parents=True)
    make_pdf(folder / "RFP.pdf", [
        "1. Scope of Work\nSpeech analytics on 100% of helpline calls with sentiment and escalation alerts.",
        "2. Eligibility Criteria\nAverage annual turnover of Rs. 3 crore in last three years.\nCERT-In empanelled bidders only."])
    r = run(settings, session)
    assert r.stats.assessed == 1 and r.stats.documents_parsed == 1 and r.stats.deep_calls == 1
    session.refresh(voice)
    assert voice.analysis_state == "done" and voice.intelligence["documents_parsed"] == 1
    doc = session.scalar(select(OpportunityDocument).where(OpportunityDocument.opportunity_id == voice.id))
    assert doc.pages == 2 and not doc.needs_manual_review
    f = session.scalar(select(EligibilityFindingRow).where(EligibilityFindingRow.opportunity_id == voice.id,
                                                           EligibilityFindingRow.criterion == "turnover"))
    assert f.value == 3 and f.evidence["document"] == "RFP.pdf" and f.evidence["page"] == 2
    assert any("RFP.pdf p.2" in e for e in voice.intelligence["evidence"])
    # a further run with no new documents does not touch it again
    assert run(settings, session).stats.assessed == 0


def test_capability_kb_change_triggers_reassessment(settings, session, tmp_path):
    write_records(settings.sources[0].path, sample_records())
    run(settings, session)
    kb_copy = tmp_path / "kb.yaml"
    kb_copy.write_text(
        Path(settings.app.capabilities_file).read_text(encoding="utf-8").replace('kb_version: "2026-10-10.1"', 'kb_version: "2026-10-11.1"'),
        encoding="utf-8",
    )
    settings.app.capabilities_file = kb_copy
    r = run(settings, session)
    assert r.stats.assessed == 6


def test_cross_source_duplicate_and_unavailable_source(settings, session, tmp_path):
    write_records(settings.sources[0].path, sample_records())
    manual = tmp_path / "manual"
    write_records(manual, [{**sample_records()[0], "platform": "GeM (manual)", "description": "From the downloaded RFP"}])
    settings.sources.append(SourceConfig(id="manual", type="manual", path=manual))
    settings.sources.append(SourceConfig(id="cppp_central", type="cppp", feed="cpppdata", verification_status="verified_working"))

    def down(req):
        if req.url.path == "/robots.txt":
            return httpx.Response(404)
        return httpx.Response(503)

    client = PoliteClient(HttpSettings(min_interval_s=0, max_retries=1, backoff_base_s=0), transport=httpx.MockTransport(down),
                          sleep=lambda s: None, resolve_dns=False)
    r = run(settings, session, http_client=client)
    assert r.run.status == "partial" and r.exit_code == 3
    assert r.stats.sources["cppp_central"]["status"] == "failed"
    assert any("cppp_central" in n for n in r.stats.coverage_notes)
    assert r.stats.duplicates == 1 and r.stats.new == 6         # good sources unaffected, no duplicate row created
    assert session.scalar(select(func.count()).select_from(Opportunity)) == 6
    wb = load_workbook(r.report_path)
    stat_rows = [row for row in wb["Scan Statistics"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert any(row[0] == "cppp_central" and row[1] == "failed" for row in stat_rows)


def test_budget_limit_enforced_and_reported(settings, session):
    write_records(settings.sources[0].path, sample_records())
    settings.budget.max_calls_per_run = 0
    r = run(settings, session)
    assert r.stats.llm_skipped_budget >= 1 and r.stats.budget_blocked_reasons
    skipped = session.scalars(select(Opportunity).where(Opportunity.analysis_state == "skipped_budget")).all()
    assert skipped and all(any("SKIPPED" in f for f in o.intelligence["flags"]) for o in skipped)
    assert llm_calls(session, r.run.id) == 0
    # next run with budget restored retries only the skipped ones
    settings.budget.max_calls_per_run = 60
    r2 = run(settings, session)
    assert r2.stats.assessed == len(skipped)


def test_llm_disabled_still_produces_report(settings, session):
    write_records(settings.sources[0].path, sample_records())
    settings.llm.mode = "disabled"
    r = run_pipeline(settings, session)
    assert r.exit_code == 0 and r.report_path.exists()
    va = session.scalar(select(Opportunity).where(Opportunity.source_tender_id == "2026_SYNSC_000101_1"))
    assert va.analysis_state == "skipped_llm_disabled" and va.priority in ("P1", "P2", "P3")


# ---------------------------------------------------------------- Excel
def test_workbook_structure_and_safety(settings, session):
    recs = sample_records()
    recs.append({"tender_id": "2026_SYNX_000199_1", "title": "=HYPERLINK(\"http://evil\",\"AI video analytics\")",
                 "organisation": "@SUM(1+1)", "closing_at": ist(12), "official_url": "https://evil.example.com/phish"})
    write_records(settings.sources[0].path, recs)
    r = run(settings, session)
    wb = load_workbook(r.report_path)
    assert wb.sheetnames == ["Executive Summary", "All Opportunities", "High Priority", "Closing Soon", "Product-Wise",
                             "Eligibility Review", "Change History", "Scan Statistics"]
    ws = wb["All Opportunities"]
    headers = [c.value for c in ws[1]]
    assert headers[:35] == REQUIRED_COLUMNS
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                assert c.data_type != "f", f"formula found in {sheet.title}!{c.coordinate}"
                if c.hyperlink:
                    assert "eprocure.gov.in" in c.hyperlink.target or "etenders.gov.in" in c.hyperlink.target or \
                        c.hyperlink.target.startswith("https://") and "evil" not in c.hyperlink.target
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    ids = [r_[0] for r_ in rows]
    assert len(ids) == len(set(ids))                               # no duplicates
    prios = [r_[17] for r_ in rows]
    assert prios == sorted(prios)                                   # P1 → P4
    titles = [r_[2] for r_ in rows]
    assert any(t.startswith("'=HYPERLINK") for t in titles)
    assert "Not stated" in [v for r_ in rows for v in r_]
    days_col = [r_[10] for r_ in rows if isinstance(r_[10], (int, float))]
    assert days_col and all(-30 < d < 60 for d in days_col)
    summary = {r_[0]: r_[1] for r_ in wb["Executive Summary"].iter_rows(min_row=3, values_only=True) if r_[0]}
    assert summary["New opportunities discovered (this run)"] == 7
    assert session.scalar(select(func.count()).select_from(ExportedReport)) == 1


# ---------------------------------------------------------------- email
def test_email_idempotent_and_dry_run(settings, session):
    write_records(settings.sources[0].path, sample_records())
    r = run(settings, session)
    rep = session.get(ExportedReport, r.report_id)
    rows = send_digest(session, settings, r.run, rep)
    assert {x.status for x in rows} == {"sent"}
    before = len(list(settings.email.outbox_dir.glob("*.eml")))
    rows2 = send_digest(session, settings, r.run, rep)              # same report + recipients → not resent
    assert [x.id for x in rows2] == [x.id for x in rows]
    assert len(list(settings.email.outbox_dir.glob("*.eml"))) == before
    dry = send_digest(session, settings, r.run, rep, dry_run=True)
    assert {x.status for x in dry} == {"dry_run"} and list((settings.email.outbox_dir / "dry_run").glob("*.eml"))


def test_email_smtp_auth_failure_recorded_not_raised(settings, session, monkeypatch):
    write_records(settings.sources[0].path, sample_records())
    r = run(settings, session)

    class BadSMTP:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def starttls(self):
            pass

        def login(self, u, p):
            raise smtplib.SMTPAuthenticationError(535, b"bad credentials")

        def send_message(self, m):
            raise AssertionError("must not send")

    monkeypatch.setattr(smtplib, "SMTP", BadSMTP)
    monkeypatch.setenv("CHANAKYA_SMTP_USERNAME", "u")
    monkeypatch.setenv("CHANAKYA_SMTP_PASSWORD", "p")
    settings.email.provider = "smtp"
    rows = send_digest(session, settings, r.run, session.get(ExportedReport, r.report_id), provider=SMTPEmailProvider(settings))
    assert {x.status for x in rows} == {"failed"} and all("authentication" in x.last_error for x in rows)
    assert session.scalar(select(func.count()).select_from(ScanRun)) == 1   # failed email did not re-run the scan


def test_email_retries_transient_errors(settings, session):
    from chanakya.reporting.email import EmailError
    write_records(settings.sources[0].path, sample_records())
    r = run(settings, session)

    class Flaky:
        n = 0

        def send(self, msg):
            Flaky.n += 1
            if Flaky.n == 1:
                raise EmailError("421 try later")
            return "<id@x>"

    settings.email.cc = []
    rows = send_digest(session, settings, r.run, None, provider=Flaky(), sleep=lambda s: None)
    assert rows[0].status == "sent" and rows[0].attempts == 2


def test_digest_escapes_untrusted_html(settings, session):
    write_records(settings.sources[0].path, [{**sample_records()[0], "title": "<script>alert(1)</script> AI video analytics on CCTV cameras"}])
    r = run(settings, session)
    d = build_digest(session, settings, r.run)
    assert "<script>" not in d.html and "&lt;script&gt;" in d.html


def test_no_recipients_is_an_error(settings, session):
    from chanakya.reporting.email import EmailError
    settings.email.to, settings.email.cc = [], []
    with pytest.raises(EmailError):
        send_digest(session, settings, None, None)


# ---------------------------------------------------------------- CLI
def test_cli_dry_run_end_to_end(settings, tmp_path):
    write_records(settings.sources[0].path, sample_records())
    cfg = tmp_path / "config" / "chanakya.yaml"
    cfg.parent.mkdir()
    data = settings.model_dump(mode="json", exclude={"config_path"})
    data["app"]["capabilities_file"] = str(settings.app.capabilities_file)
    cfg.write_text(yaml.safe_dump(data))
    runner = CliRunner()
    res = runner.invoke(cli_app, ["--config", str(cfg), "run", "--dry-run", "--llm-mode", "fake"])
    assert res.exit_code == 0, res.output
    assert "Run run_" in res.output and "report:" in res.output
    assert list((settings.email.outbox_dir / "dry_run").glob("*.eml"))
    assert not list(settings.email.outbox_dir.glob("*.eml"))        # nothing "sent" in dry-run
    res2 = runner.invoke(cli_app, ["--config", str(cfg), "status"])
    assert res2.exit_code == 0 and "Recent runs" in res2.output
    res3 = runner.invoke(cli_app, ["--config", str(cfg), "costs", "--last-days", "7"])
    assert res3.exit_code == 0 and "Total last 7 days" in res3.output
    res4 = runner.invoke(cli_app, ["--config", str(cfg), "report"])
    assert res4.exit_code == 0
    res5 = runner.invoke(cli_app, ["--config", str(cfg), "capabilities", "validate"])
    assert res5.exit_code == 0 and "21 capabilities" in res5.output


def test_benchmark_funnel_is_cheaper_without_losing_recall(example_settings):
    from chanakya.benchmark import run_benchmark
    from .conftest import FIXTURES
    res = run_benchmark(example_settings, FIXTURES / "benchmark" / "labeled_tenders.yaml")
    a = {x.name: x for x in res.approaches}
    assert a["chanakya"].cost_usd < 0.2 * a["llm_all"].cost_usd
    assert a["chanakya"].recall >= a["keyword_llm"].recall
    assert a["chanakya"].calls_fast + a["chanakya"].calls_deep < a["keyword_llm"].calls_deep


def test_page_cap_keeps_scanning_and_does_not_skip_the_rest_of_the_window(settings, session, monkeypatch):
    from chanakya.connectors.base import TenderSourceConnector
    from chanakya.db import Source
    from chanakya.models import ListingPage, TenderRecord
    from chanakya.pipeline import runner as runner_mod
    from chanakya.pipeline.runner import Pipeline, RunStats
    from chanakya.timeutil import now_ist, now_utc

    monkeypatch.setattr(runner_mod, "_PAGE_SAFETY", 3)

    class Pager(TenderSourceConnector):
        def fetch_listing(self, url, page_no):
            n = int(url or "0")
            rec = TenderRecord(
                source_id=self.source_id, platform="Pager", source_record_key=f"k{n}",
                title=f"Notice {n}", published_at=now_ist(), official_url="https://example.test/t",
                fetched_at=now_utc(),
            )
            return ListingPage(records=[rec], next_url=str(n + 1), page_no=page_no)

    monkeypatch.setattr(runner_mod, "build_connector", lambda cfg, client=None: Pager(cfg, client))
    settings.sources = [SourceConfig(id="pager", type="manual", max_pages=2)]
    pipe = Pipeline(settings, session)
    stats = RunStats()
    pipe.scan(pipe.start_run(trigger="test"), stats)
    src = session.get(Source, "pager")
    assert stats.sources["pager"]["pages"] >= 3
    assert src.high_water_mark is None
    assert any("coverage may be partial" in n for n in stats.coverage_notes)

    class Done(TenderSourceConnector):
        def fetch_listing(self, url, page_no):
            rec = TenderRecord(
                source_id=self.source_id, platform="Pager", source_record_key="done",
                title="Done notice", published_at=now_ist(), official_url="https://example.test/t",
                fetched_at=now_utc(),
            )
            return ListingPage(records=[rec], next_url=None, page_no=page_no)

    monkeypatch.setattr(runner_mod, "build_connector", lambda cfg, client=None: Done(cfg, client))
    pipe.scan(pipe.start_run(trigger="test"), RunStats())
    session.refresh(src)
    assert src.high_water_mark is not None


def test_email_delivery_rows_have_unique_keys(settings, session):
    write_records(settings.sources[0].path, sample_records())
    run(settings, session, email=True)
    keys = [x.idempotency_key for x in session.scalars(select(EmailDelivery))]
    assert len(keys) == len(set(keys)) == 2


def test_notices_are_assessed_while_the_scan_is_still_running(settings, session, monkeypatch):
    """A pass that dies before its closing assessment must not leave the book unscored."""
    from chanakya.pipeline.runner import Pipeline

    def dies(self, *a, **kw):
        raise RuntimeError("process ended before the closing assessment")

    monkeypatch.setattr(Pipeline, "assess", dies)
    write_records(settings.sources[0].path, sample_records())
    res = run(settings, session, export=False)
    assert res.run.status == "failed" and res.stats.assessed == 6
    opps = {o.source_tender_id: o for o in session.scalars(select(Opportunity))}
    assert all(o.filter_decision and o.priority for o in opps.values())
    assert opps["2026_SYNOLD_000105_1"].status == "expired" and opps["2026_SYNOLD_000105_1"].priority == "P4"


def test_per_run_call_limit_holds_across_scan_checkpoints(settings, session):
    settings.budget.max_calls_per_run = 1
    write_records(settings.sources[0].path, sample_records())
    r = run(settings, session, export=False)
    assert llm_calls(session, r.run.id) <= 1

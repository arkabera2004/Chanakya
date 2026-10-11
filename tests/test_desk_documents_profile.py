"""The desk takes tender documents for one notice, and the company profile eligibility is checked against.

Every test runs on a temporary database, a temporary inbox and the simulated
model. Nothing here touches ``data/``, the real config or the network. The
documents are SYNTHETIC: they are built in memory for these tests.
"""

import io
import os
import time
import zipfile

import pytest
import yaml
from fastapi.testclient import TestClient

from chanakya import server
from chanakya.app import open_session, run_pipeline
from chanakya.config import apply_profile_overlay, load_settings, profile_overlay_path
from chanakya.desklock import acquire, lock_path, release
from chanakya.server import create_app

from .conftest import ROOT, sample_records, write_records

PDF = b"%PDF-1.4\n% SYNTHETIC test document\n" + b"0" * 200

ELIGIBILITY_LINES = [
    "Scope of Work",
    "SYNTHETIC tender document. AI/ML based automated scrutiny of building plan applications, "
    "with document understanding, rule checks against the building bye-laws and a reviewer console.",
    "Eligibility Criteria",
    "The bidder must have an average annual turnover of Rs. 5 crore in the last three financial years.",
    "The bidder must be ISO 27001 certified.",
]


def docx_bytes(lines: list[str] = ELIGIBILITY_LINES) -> bytes:
    import docx

    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["Item", "Qty"])
    book.active.append(["Plan scrutiny licence (SYNTHETIC)", 1])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def zip_bytes(members: dict[str, bytes], compression: int = zipfile.ZIP_STORED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.fixture
def desk(settings, monkeypatch):
    """A scored book of synthetic notices. The desk reloads its config on every request,
    which is when a saved profile is laid over the YAML, so the patched loader does the same."""
    write_records(settings.sources[0].path, sample_records())
    settings.http.max_download_mb = 1
    with open_session(settings) as session:
        run_pipeline(settings, session, email=False, export=False)
    monkeypatch.setattr(server, "desk_settings", lambda: apply_profile_overlay(settings.model_copy(deep=True)))
    monkeypatch.setattr(server, "_job", None)
    return settings


@pytest.fixture
def client(desk):
    with TestClient(create_app()) as test_client:
        yield test_client


def _lead(client: TestClient, text: str) -> dict:
    rows = client.get("/api/opportunities", params={"q": text, "include_rejected": True}).json()["opportunities"]
    assert len(rows) == 1, rows
    return rows[0]


def _upload(client: TestClient, lead_id: str, name: str, data):
    return client.post(f"/api/opportunities/{lead_id}/documents", params={"filename": name}, content=data)


def _folder(settings, lead_id: str):
    return settings.app.inbox_dir / "documents" / lead_id


def _wait_for_job(client: TestClient, tries: int = 240, pause: float = 0.05) -> dict:
    job = None
    for _ in range(tries):
        job = client.get("/api/job").json()["job"]
        if job and job["status"] != "running":
            return job
        time.sleep(pause)
    raise AssertionError(f"job did not finish: {job}")


def _nothing_left_behind(settings) -> bool:
    incoming = settings.app.inbox_dir / ".uploads"
    return not incoming.exists() or not list(incoming.iterdir())


# ------------------------------------------------------------------ documents


def test_upload_stores_the_file_where_a_pass_looks(client, desk):
    lead = _lead(client, "building plan")
    empty = client.get(f"/api/opportunities/{lead['id']}/documents").json()
    assert empty["documents"] == []
    assert empty["max_documents"] == desk.budget.max_documents_per_tender and empty["max_mb"] == 1

    payload = docx_bytes()  # built once: a .docx carries its creation time, so two builds differ
    sent = _upload(client, lead["id"], "NIT and eligibility.docx", payload)
    assert sent.status_code == 200, sent.text
    body = sent.json()
    assert body["stored"] == {
        "name": "NIT and eligibility.docx", "type": "docx", "size_bytes": len(payload),
        "replaced": False, "archive_members": [],
    }
    stored = _folder(desk, lead["id"]) / "NIT and eligibility.docx"
    assert stored.read_bytes() == payload
    assert _nothing_left_behind(desk)

    [row] = client.get(f"/api/opportunities/{lead['id']}/documents").json()["documents"]
    assert row["name"] == "NIT and eligibility.docx" and row["type"] == "docx"
    assert row["origin"] == "inbox" and row["removable"] and row["within_limit"]
    assert row["read"] is False and row["pages"] is None  # stored, not read: reading is the pass's job

    # The pipeline finds the same file, by the same rule, with no desk in between.
    from chanakya.db import Opportunity
    from chanakya.pipeline.runner import documents_for

    with open_session(desk) as session:
        assert documents_for(desk, session.get(Opportunity, lead["id"])) == [stored]

    assert _upload(client, "CHK-NOSUCHNOTICE", "a.pdf", PDF).status_code == 404
    assert client.get("/api/opportunities/CHK-NOSUCHNOTICE/documents").status_code == 404


def test_the_type_comes_from_the_bytes_not_from_the_name(client, desk):
    lead = _lead(client, "building plan")
    # A PDF that calls itself a text file is stored as the PDF it is.
    renamed = _upload(client, lead["id"], "tender.txt", PDF)
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["stored"]["name"] == "tender.txt.pdf"
    assert _upload(client, lead["id"], "boq.xlsx", xlsx_bytes()).json()["stored"]["type"] == "xlsx"

    refused = {
        "notes.pdf": b"plain text that only claims to be a PDF",
        "page.pdf": b"<html><body><table><tr><td>not a document</td></tr></table></body></html>",
        "tool.pdf": b"MZ\x90\x00\x03\x00\x00\x00\xff\xfe\xfd binary",
        "broken.docx": b"PK\x03\x04 not really an archive",
    }
    for name, data in refused.items():
        answer = _upload(client, lead["id"], name, data)
        assert answer.status_code == 415, (name, answer.text)
        assert "PDF, DOCX, XLSX or ZIP" in answer.json()["detail"]
    legacy = _upload(client, lead["id"], "old.doc", b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    assert legacy.status_code == 415 and "Legacy" in legacy.json()["detail"]
    assert _upload(client, lead["id"], "empty.pdf", b"").status_code == 400

    assert sorted(p.name for p in _folder(desk, lead["id"]).iterdir()) == ["boq.xlsx", "tender.txt.pdf"]
    assert _nothing_left_behind(desk)


def test_an_oversize_upload_is_refused_and_not_kept(client, desk):
    lead = _lead(client, "building plan")
    too_big = PDF + b"0" * (1024 * 1024)
    declared = _upload(client, lead["id"], "big.pdf", too_big)
    assert declared.status_code == 413
    assert "1 MB" in declared.json()["detail"]
    # No Content-Length to trust: the cap still holds while the body is read.
    streamed = _upload(client, lead["id"], "big.pdf", iter([too_big[:4096], too_big[4096:]]))
    assert streamed.status_code == 413
    assert not _folder(desk, lead["id"]).exists()
    assert _nothing_left_behind(desk)
    just_inside = PDF + b"0" * (1024 * 1024 - len(PDF))
    assert _upload(client, lead["id"], "fits.pdf", just_inside).status_code == 200


def test_an_unsafe_filename_cannot_leave_the_notice_folder(client, desk, tmp_path):
    lead = _lead(client, "building plan")
    folder = _folder(desk, lead["id"])
    before = {p for p in tmp_path.rglob("*") if p.is_file()}

    for sent, kept in [
        ("../../../escape.pdf", "escape.pdf"),
        ("C:\\Windows\\System32\\evil.pdf", "evil.pdf"),
        ("CON.pdf", "_CON.pdf"),
    ]:
        answer = _upload(client, lead["id"], sent, PDF)
        assert answer.status_code == 200, (sent, answer.text)
        assert answer.json()["stored"]["name"] == kept
        client.delete(f"/api/opportunities/{lead['id']}/documents/{kept}")
    odd = _upload(client, lead["id"], "bad\x00name<script>.pdf", PDF)
    assert odd.status_code == 200 and odd.json()["stored"]["name"] == "badname_script_.pdf"

    for nameless in ("", "..", "../", "///", " . "):
        assert _upload(client, lead["id"], nameless, PDF).status_code == 400, nameless

    after = {p for p in tmp_path.rglob("*") if p.is_file()}
    assert after - before == {folder / "badname_script_.pdf"}

    # Removing is held to the same folder.
    (desk.app.inbox_dir / "documents" / "secret.pdf").write_bytes(PDF)
    for name in ("..%2Fsecret.pdf", "..\\secret.pdf", "nosuch.pdf"):
        assert client.delete(f"/api/opportunities/{lead['id']}/documents/{name}").status_code in (400, 404)
    assert (desk.app.inbox_dir / "documents" / "secret.pdf").exists()


def test_a_portal_reference_cannot_name_a_folder_outside_the_inbox(desk, tmp_path):
    """The Tender ID and reference are portal text. They pick a folder, so they are sanitised first."""
    from chanakya.db import Opportunity
    from chanakya.pipeline.runner import document_folders, documents_for

    outside = tmp_path / "private"
    outside.mkdir()
    (outside / "payroll.pdf").write_bytes(PDF)
    inbox = desk.app.inbox_dir / "documents"
    notice = Opportunity(id="CHK-TRAVERSE", source_tender_id="../../private", ref_no="..")
    folders = document_folders(desk, notice)
    assert all(folder.parent == inbox for folder in folders)
    assert all(folder.resolve().parent == inbox.resolve() for folder in folders)
    assert documents_for(desk, notice) == []


def test_archives_are_opened_the_guarded_way(client, desk, tmp_path):
    lead = _lead(client, "building plan")
    good = zip_bytes({"tender/NIT.docx": docx_bytes(), "tender/boq.xlsx": xlsx_bytes(), "tender/readme.txt": b"hello"})
    kept = _upload(client, lead["id"], "bundle.zip", good)
    assert kept.status_code == 200, kept.text
    assert kept.json()["stored"]["type"] == "zip"
    assert sorted(kept.json()["stored"]["archive_members"]) == ["NIT.docx", "boq.xlsx"]
    # The archive is kept as one document. It is not unpacked into the inbox.
    assert [p.name for p in _folder(desk, lead["id"]).iterdir()] == ["bundle.zip"]

    refused = {
        "traversal.zip": zip_bytes({"../../escape.pdf": PDF}),
        "absolute.zip": zip_bytes({"/etc/escape.pdf": PDF}),
        "bomb.zip": zip_bytes({"zeros.pdf": b"%PDF" + b"\0" * (4 * 1024 * 1024)}, zipfile.ZIP_DEFLATED),
        "nothing-to-read.zip": zip_bytes({"readme.txt": b"hello", "run.bat": b"echo hi"}),
    }
    for name, data in refused.items():
        answer = _upload(client, lead["id"], name, data)
        assert answer.status_code == 400, (name, answer.text)
        assert "archive" in answer.json()["detail"]
    assert "zip bomb" in _upload(client, lead["id"], "bomb.zip", refused["bomb.zip"]).json()["detail"]
    assert [p.name for p in _folder(desk, lead["id"]).iterdir()] == ["bundle.zip"]
    assert not list(tmp_path.rglob("escape.pdf"))
    assert _nothing_left_behind(desk)


def test_a_notice_takes_only_as_many_documents_as_a_pass_reads(client, desk):
    lead = _lead(client, "building plan")
    limit = desk.budget.max_documents_per_tender
    for index in range(limit):
        assert _upload(client, lead["id"], f"part-{index}.pdf", PDF).status_code == 200
    over = _upload(client, lead["id"], "one-too-many.pdf", PDF)
    assert over.status_code == 400
    assert "max_documents_per_tender" in over.json()["detail"]
    assert not (_folder(desk, lead["id"]) / "one-too-many.pdf").exists()

    # A new copy of a document already there replaces it and is not counted twice.
    again = _upload(client, lead["id"], "part-0.pdf", PDF + b"revised")
    assert again.status_code == 200 and again.json()["stored"]["replaced"] is True

    # A document dropped in by hand, under the Tender ID, counts too: a pass would read it.
    by_hand = desk.app.inbox_dir / "documents" / lead["tender_id"]
    by_hand.mkdir(parents=True)
    (by_hand / "from-cli.pdf").write_bytes(PDF)
    rows = client.get(f"/api/opportunities/{lead['id']}/documents").json()["documents"]
    assert len(rows) == limit + 1
    assert [row["within_limit"] for row in rows].count(True) == limit

    removed = client.delete(f"/api/opportunities/{lead['id']}/documents/from-cli.pdf", params={"folder": lead["tender_id"]})
    assert removed.status_code == 200 and len(removed.json()["documents"]) == limit
    assert client.delete(f"/api/opportunities/{lead['id']}/documents/part-1.pdf").status_code == 200
    assert _upload(client, lead["id"], "one-too-many.pdf", PDF).status_code == 200


def test_a_watch_holding_the_database_gets_a_409(client, desk):
    lead = _lead(client, "building plan")
    assert _upload(client, lead["id"], "kept.pdf", PDF).status_code == 200
    path = lock_path(desk)
    acquire(path, {"pid": os.getpid(), "desk": "live", "mode": "watch", "started_at": "2026-10-10T00:00:00Z", "pause_s": 600})
    try:
        refused = _upload(client, lead["id"], "during-watch.pdf", PDF)
        assert refused.status_code == 409
        assert "watch" in refused.json()["detail"].lower()
        assert client.delete(f"/api/opportunities/{lead['id']}/documents/kept.pdf").status_code == 409
        assert client.post(f"/api/opportunities/{lead['id']}/analyze").status_code == 409
        # Reading is not writing.
        assert client.get(f"/api/opportunities/{lead['id']}/documents").status_code == 200
    finally:
        release(path, os.getpid())
    assert [p.name for p in _folder(desk, lead["id"]).iterdir()] == ["kept.pdf"]
    assert _nothing_left_behind(desk)
    assert _upload(client, lead["id"], "after-watch.pdf", PDF).status_code == 200


def test_analysing_one_notice_reads_its_documents_and_leaves_the_rest(client, desk):
    lead = _lead(client, "building plan")
    other = _lead(client, "video analytics")
    before = client.get(f"/api/opportunities/{lead['id']}").json()
    assert before["findings"] == []
    assert before["analysis_state"] == "deferred_no_documents"  # the listing alone is too thin for a deep read

    assert _upload(client, lead["id"], "NIT.docx", docx_bytes()).status_code == 200
    started = client.post(f"/api/opportunities/{lead['id']}/analyze")
    assert started.status_code == 200, started.text
    assert started.json()["job"]["opportunity_id"] == lead["id"]
    job = _wait_for_job(client)
    assert job["status"] == "ok", job
    assert job["stats"]["trigger"] == "notice"
    assert job["stats"]["assessed"] == 1 and job["stats"]["deep_calls"] == 1
    assert [event["opportunity_id"] for event in job["events"] if event["kind"] == "judge"] == [lead["id"], lead["id"]]

    after = client.get(f"/api/opportunities/{lead['id']}").json()
    assert after["analysis_state"] == "done"
    assert after["confidence_basis"] == "AI deep analysis"
    assert "SIMULATED AI OUTPUT (fake provider)" in after["flags"]  # the simulated model says it is simulated
    found = {(finding["criterion"], finding["unit"]): finding for finding in after["findings"]}
    assert found[("turnover", "INR crore")]["value"] == 5.0
    assert ("certification", "ISO 27001") in found
    # Nothing is known about Wayam yet, so nothing is claimed.
    assert {finding["assessment"] for finding in found.values()} == {"unknown"}
    assert after["eligibility"].startswith("Unknown")
    assert any("NIT.docx" in line for line in after["evidence"])

    [row] = client.get(f"/api/opportunities/{lead['id']}/documents").json()["documents"]
    assert row["read"] is True and row["pages"] == 1 and row["error"] is None

    untouched = client.get(f"/api/opportunities/{other['id']}").json()
    assert untouched["analysis_state"] == other["analysis_state"] and untouched["score"] == other["score"]
    assert client.post("/api/opportunities/CHK-NOSUCHNOTICE/analyze").status_code == 404


def test_analyse_one_stays_inside_the_model_budget(client, desk):
    """One notice is still a pass: the budget guard decides whether the model is called."""
    lead = _lead(client, "building plan")
    desk.budget.max_calls_per_run = 0
    assert _upload(client, lead["id"], "NIT.docx", docx_bytes()).status_code == 200
    assert client.post(f"/api/opportunities/{lead['id']}/analyze").status_code == 200
    job = _wait_for_job(client)
    assert job["stats"]["deep_calls"] == 0 and job["stats"]["llm_skipped_budget"] >= 1
    after = client.get(f"/api/opportunities/{lead['id']}").json()
    assert after["analysis_state"] == "skipped_budget"
    # The clauses are read by rule, so the findings arrive without the model.
    assert {finding["criterion"] for finding in after["findings"]} >= {"turnover", "certification"}


# -------------------------------------------------------------------- profile


EMPTY = {
    "annual_turnover_inr_crore": None, "average_turnover_3y_inr_crore": None, "years_in_operation": None,
    "certifications": [], "msme_registered": None, "dpiit_startup": None,
    "past_government_projects": None, "cert_in_empanelled": None,
}


def test_the_profile_starts_not_stated_and_survives_a_reload(client, desk):
    first = client.get("/api/profile").json()
    assert first["saved"] is False and first["saved_at"] is None
    assert {key: first["profile"][key] for key in EMPTY} == EMPTY  # nothing is filled in for anyone

    filled = {
        "average_turnover_3y_inr_crore": 12.5, "years_in_operation": 6, "certifications": ["ISO 27001", " iso 27001 ", "", "CMMI 3"],
        "msme_registered": True, "cert_in_empanelled": False, "past_government_projects": 2,
    }
    saved = client.put("/api/profile", json=filled)
    assert saved.status_code == 200, saved.text
    profile = saved.json()["profile"]
    assert profile["average_turnover_3y_inr_crore"] == 12.5 and profile["years_in_operation"] == 6
    assert profile["certifications"] == ["ISO 27001", "CMMI 3"]
    assert profile["msme_registered"] is True and profile["cert_in_empanelled"] is False
    assert profile["annual_turnover_inr_crore"] is None and profile["dpiit_startup"] is None  # untouched stays not stated
    assert profile["name"] == desk.company_profile.name
    assert saved.json()["saved"] is True and saved.json()["saved_at"]

    path = profile_overlay_path(desk)
    assert path == desk.app.data_dir / "chanakya.example.profile.yaml" and path.is_file()
    assert client.get("/api/profile").json()["profile"] == profile

    # Clearing a field makes it not stated again.
    cleared = client.put("/api/profile", json={**filled, "average_turnover_3y_inr_crore": None, "certifications": []})
    assert cleared.json()["profile"]["average_turnover_3y_inr_crore"] is None
    assert cleared.json()["profile"]["certifications"] == []


def test_the_profile_refuses_values_it_cannot_stand_behind(client, desk):
    for bad in (
        {"average_turnover_3y_inr_crore": -1},
        {"years_in_operation": "about six"},
        {"msme_registered": "yes"},
        {"past_government_projects": 1.5},
        {"net_worth_inr_crore": 3},  # not a field: a typo must not be kept silently
        {"certifications": ["x" * 81]},
    ):
        answer = client.put("/api/profile", json=bad)
        assert answer.status_code == 400, (bad, answer.text)
        assert isinstance(answer.json()["detail"], str)
    assert not profile_overlay_path(desk).exists()
    assert client.get("/api/profile").json()["saved"] is False


def test_a_saved_profile_is_laid_over_the_yaml_without_rewriting_it(tmp_path):
    """The real load path: the YAML is the base, keeps its comments, and the saved file wins."""
    config = tmp_path / "config" / "chanakya.yaml"
    config.parent.mkdir()
    text = (
        "# the owner's notes stay here\n"
        "company_profile:\n"
        "  name: Wayam AI\n"
        "  years_in_operation: 4          # from the incorporation certificate\n"
        "  annual_turnover_inr_crore: null\n"
    )
    config.write_text(text, encoding="utf-8")
    base = load_settings(config)
    assert base.company_profile.years_in_operation == 4
    overlay = profile_overlay_path(base)
    assert overlay == tmp_path / "data" / "chanakya.profile.yaml"

    overlay.parent.mkdir(parents=True, exist_ok=True)
    overlay.write_text(yaml.safe_dump({"company_profile": {
        "annual_turnover_inr_crore": 8.0, "years_in_operation": None, "certifications": ["ISO 9001"],
    }}), encoding="utf-8")
    merged = load_settings(config).company_profile
    assert merged.annual_turnover_inr_crore == 8.0 and merged.certifications == ["ISO 9001"]
    assert merged.years_in_operation is None  # a saved null is "not stated", over the YAML's value
    assert merged.name == "Wayam AI"
    assert config.read_text(encoding="utf-8") == text

    # Another config in the same project does not inherit these figures.
    other = tmp_path / "config" / "chanakya.demo.yaml"
    other.write_text("company_profile:\n  name: Wayam AI\n", encoding="utf-8")
    assert load_settings(other).company_profile.annual_turnover_inr_crore is None

    overlay.write_text("company_profile:\n  turnover: 8\n", encoding="utf-8")
    with pytest.raises(ValueError):
        load_settings(config)


def test_the_shipped_configs_state_no_credentials():
    """Wayam's credentials are never invented: the example and demo configs leave every one empty."""
    for name in ("chanakya.example.yaml", "chanakya.demo.yaml"):
        raw = yaml.safe_load((ROOT / "config" / name).read_text(encoding="utf-8")) or {}
        stated = {key: value for key, value in (raw.get("company_profile") or {}).items() if key != "name"}
        assert all(value in (None, []) for value in stated.values()), (name, stated)


def test_a_saved_profile_changes_eligibility_on_the_next_assessment(client, desk):
    lead = _lead(client, "video analytics")  # asks for Rs. 10 crore average turnover and ISO 27001
    before = client.get(f"/api/opportunities/{lead['id']}").json()
    assert before["eligibility"].startswith("Unknown")
    gating = [f for f in before["findings"] if f["criterion"] in ("turnover", "certification")]
    assert len(gating) == 2 and {f["assessment"] for f in gating} == {"unknown"}

    def reassess(profile: dict) -> dict:
        assert client.put("/api/profile", json=profile).status_code == 200
        started = client.post("/api/analyze", json={"scope": "all"})
        assert started.status_code == 200, started.text
        assert _wait_for_job(client)["status"] == "ok"
        return client.get(f"/api/opportunities/{lead['id']}").json()

    met = reassess({"average_turnover_3y_inr_crore": 12, "certifications": ["ISO 27001"]})
    assert {f["assessment"] for f in met["findings"] if f["criterion"] in ("turnover", "certification")} == {"meets"}
    assert met["eligibility"] != before["eligibility"]
    assert met["eligibility"] in ("Verified eligible", "Potentially eligible")

    short = reassess({"average_turnover_3y_inr_crore": 2, "certifications": ["ISO 27001"]})
    assert short["eligibility"] == "Likely ineligible"
    assert [f["assessment"] for f in short["findings"] if f["criterion"] == "turnover"] == ["does_not_meet"]

    # Taking the figures away again returns the notice to unknown. Nothing is remembered as met.
    assert reassess({})["eligibility"].startswith("Unknown")
